import pdfplumber
from docx import Document
import re
from typing import Optional, Tuple
import io
from app.core.config import settings


class FileParser:
    @staticmethod
    def parse_pdf(file_content: bytes) -> str:
        """Парсинг PDF файла"""
        text = ""
        with pdfplumber.open(io.BytesIO(file_content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text.strip()

    @staticmethod
    def parse_docx(file_content: bytes) -> str:
        """Парсинг DOCX файла"""
        doc = Document(io.BytesIO(file_content))
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)

        # Также проверяем таблицы
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)

        return "\n".join(text)

    @staticmethod
    def parse_text(file_content: bytes, encoding: str = 'utf-8') -> str:
        """Парсинг текстового файла"""
        try:
            return file_content.decode(encoding)
        except UnicodeDecodeError:
            # Пробуем другие кодировки
            for enc in ['latin-1', 'cp1251', 'cp1252']:
                try:
                    return file_content.decode(enc)
                except UnicodeDecodeError:
                    continue
            raise ValueError("Не удалось декодировать файл")

    @staticmethod
    def parse_file(file_content: bytes, file_extension: str) -> Tuple[str, Optional[str]]:
        """Основной метод парсинга файлов"""
        if file_extension.lower() == '.pdf':
            return FileParser.parse_pdf(file_content), 'pdf'
        elif file_extension.lower() in ['.docx', '.doc']:
            return FileParser.parse_docx(file_content), 'docx'
        elif file_extension.lower() == '.txt':
            return FileParser.parse_text(file_content), 'txt'
        else:
            raise ValueError(f"Неподдерживаемый формат файла: {file_extension}")

    @staticmethod
    def extract_sections(text: str) -> dict:
        """Извлечение секций из текста резюме"""
        sections = {
            'experience': '',
            'education': '',
            'skills': '',
            'summary': ''
        }

        # Паттерны для поиска секций
        patterns = {
            'experience': r'(?i)(experience|work\s*experience|employment\s*history|professional\s*experience)[:\s]*\n(.*?)(?=\n\s*\n|\n[A-Z][a-z]+\s*:|$)',
            'education': r'(?i)(education|academic\s*background|qualifications)[:\s]*\n(.*?)(?=\n\s*\n|\n[A-Z][a-z]+\s*:|$)',
            'skills': r'(?i)(skills|technical\s*skills|competencies)[:\s]*\n(.*?)(?=\n\s*\n|\n[A-Z][a-z]+\s*:|$)',
            'summary': r'(?i)(summary|profile|objective)[:\s]*\n(.*?)(?=\n\s*\n|\n[A-Z][a-z]+\s*:|$)'
        }

        for section, pattern in patterns.items():
            match = re.search(pattern, text, re.DOTALL | re.IGNORECASE)
            if match:
                sections[section] = match.group(2).strip()

        return sections